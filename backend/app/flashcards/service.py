from datetime import datetime

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app import models
from app.flashcard_sets.service import find_one_or_fail as find_set_or_fail, increment_card_count
from app.flashcards.schemas import CreateCardRequest, UpdateCardRequest, ImportCardsRequest


def find_all_by_user(db: Session, user_id: int, page: int, page_size: int) -> tuple[list[models.Flashcard], int]:
    from sqlalchemy import func
    # Subquery: pick the smallest id per unique (front, back) pair
    unique_ids_sq = (
        db.query(func.min(models.Flashcard.id).label("id"))
        .join(models.FlashcardSet, models.Flashcard.set_id == models.FlashcardSet.id)
        .filter(models.FlashcardSet.user_id == user_id)
        .group_by(models.Flashcard.front, models.Flashcard.back)
        .subquery()
    )
    query = (
        db.query(models.Flashcard)
        .join(unique_ids_sq, models.Flashcard.id == unique_ids_sq.c.id)
        .order_by(models.Flashcard.id.desc())
    )
    total = query.count()
    cards = query.offset((page - 1) * page_size).limit(page_size).all()
    return cards, total


def find_by_set(db: Session, set_id: int, user_id: int) -> list[models.Flashcard]:
    find_set_or_fail(db, set_id, user_id)  # ownership check
    return (
        db.query(models.Flashcard)
        .filter(models.Flashcard.set_id == set_id)
        .order_by(models.Flashcard.next_review_at.asc())
        .all()
    )


def find_one_or_fail(db: Session, card_id: int, user_id: int) -> models.Flashcard:
    card = db.query(models.Flashcard).filter(models.Flashcard.id == card_id).first()
    if not card:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Card not found")
    # verify ownership via set
    from app.flashcard_sets.service import find_one_or_fail as set_find
    set_find(db, card.set_id, user_id)
    return card


def create(db: Session, user_id: int, dto: CreateCardRequest) -> models.Flashcard:
    from app.flashcard_sets.service import find_one_or_fail as set_find
    set_find(db, dto.set_id, user_id)  # ownership check
    card = models.Flashcard(
        set_id=dto.set_id,
        front=dto.front,
        back=dto.back,
        next_review_at=datetime.utcnow(),
    )
    db.add(card)
    db.flush()
    increment_card_count(db, dto.set_id, 1)
    db.commit()
    db.refresh(card)
    return card


def update(db: Session, card_id: int, user_id: int, dto: UpdateCardRequest) -> models.Flashcard:
    card = find_one_or_fail(db, card_id, user_id)
    if dto.front is not None:
        card.front = dto.front
    if dto.back is not None:
        card.back = dto.back
    db.commit()
    db.refresh(card)
    return card


def delete(db: Session, card_id: int, user_id: int) -> None:
    card = find_one_or_fail(db, card_id, user_id)
    set_id = card.set_id
    db.delete(card)
    db.flush()
    increment_card_count(db, set_id, -1)
    db.commit()


def bulk_import(db: Session, user_id: int, dto: ImportCardsRequest) -> list[models.Flashcard]:
    from app.flashcard_sets.service import find_one_or_fail as set_find
    set_find(db, dto.set_id, user_id)
    now = datetime.utcnow()
    cards = [
        models.Flashcard(
            set_id=dto.set_id,
            front=item.front,
            back=item.back,
            next_review_at=now,
        )
        for item in dto.cards
    ]
    db.add_all(cards)
    db.flush()
    increment_card_count(db, dto.set_id, len(cards))
    db.commit()
    for c in cards:
        db.refresh(c)
    return cards


def import_from_excel(db: Session, user_id: int, file_bytes: bytes, filename: str,
                      set_id: int | None, set_name: str | None) -> dict:
    from app.flashcard_sets.service import create as create_set, find_one_or_fail as set_find
    from app.flashcard_sets.schemas import CreateSetRequest
    import io

    # Parse Excel file
    rows: list[tuple[str, str]] = []
    lower_name = filename.lower()

    if lower_name.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
        ws = wb.active
        header = [str(cell.value or "").strip().lower() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        if "russian" not in header or "english" not in header:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                                detail="File must have 'english' and 'russian' columns")
        ru_idx = header.index("russian")
        en_idx = header.index("english")
        for row in ws.iter_rows(min_row=2, values_only=True):
            vals = list(row)
            ru = str(vals[ru_idx] or "").strip() if ru_idx < len(vals) else ""
            en = str(vals[en_idx] or "").strip() if en_idx < len(vals) else ""
            if ru and en:
                rows.append((ru, en))
        wb.close()

    elif lower_name.endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(file_contents=file_bytes)
        sheet = book.sheet_by_index(0)
        header = [str(sheet.cell_value(0, c)).strip().lower() for c in range(sheet.ncols)]
        if "russian" not in header or "english" not in header:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                                detail="File must have 'english' and 'russian' columns")
        ru_idx = header.index("russian")
        en_idx = header.index("english")
        for r in range(1, sheet.nrows):
            ru = str(sheet.cell_value(r, ru_idx)).strip()
            en = str(sheet.cell_value(r, en_idx)).strip()
            if ru and en:
                rows.append((ru, en))
    else:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                            detail="Unsupported file format. Use .xls or .xlsx")

    if not rows:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="No valid rows found in file")

    # Resolve or create set
    if set_id:
        set_find(db, set_id, user_id)
        target_set_id = set_id
    elif set_name:
        language = db.query(models.Language).filter(models.Language.code == "en").first()
        if not language:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                                detail="English language not found in database")
        new_set = create_set(db, user_id, CreateSetRequest(
            name=set_name, language_id=language.id,
        ))
        target_set_id = new_set.id
    else:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                            detail="Provide set_id or set_name")

    # Create cards
    now = datetime.utcnow()
    cards = [
        models.Flashcard(set_id=target_set_id, front=ru, back=en, next_review_at=now)
        for ru, en in rows
    ]
    db.add_all(cards)
    db.flush()
    increment_card_count(db, target_set_id, len(cards))
    db.commit()

    return {"imported": len(cards), "skipped": 0, "set_id": target_set_id}


def find_due_cards(db: Session, set_id: int, limit: int) -> list[models.Flashcard]:
    now = datetime.utcnow()
    return (
        db.query(models.Flashcard)
        .filter(
            models.Flashcard.set_id == set_id,
            models.Flashcard.next_review_at <= now,
        )
        .order_by(models.Flashcard.next_review_at.asc())
        .limit(limit)
        .all()
    )


def find_oldest_cards(db: Session, set_id: int, limit: int) -> list[models.Flashcard]:
    return (
        db.query(models.Flashcard)
        .filter(models.Flashcard.set_id == set_id)
        .order_by(models.Flashcard.next_review_at.asc())
        .limit(limit)
        .all()
    )
